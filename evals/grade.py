"""Automatické hodnocení výstupů testů skillu protokol-mereni.

    python grade.py <run_dir> <eval_name>
run_dir = .../<eval>/<with_skill|without_skill>  -> zapíše grading.json
"""
import glob
import json
import os
import re
import sys

from docx import Document
from docx.oxml.ns import qn

POVINNE = ["Teoretický úvod", "Schéma zapojení", "Použité přístroje", "Postup měření",
           "Tabulky", "Příklad výpočtu", "Závěr"]
ZAKAZ = ["dle očekávání", "podle očekávání", "odpovídají teoretickým předpokladům"]
MIN_GRAFU = {"stabilizatory-cv16": 4, "feromagnetika-cv11": 2, "usmernovac-cv14-chybejici-udaje": 2}


def h1s(d):
    return [p.text.strip() for p in d.paragraphs if p.style.name.startswith("Heading 1") or p.style.name == "Title"]


def section_texts(d, name):
    ps = d.paragraphs
    try:
        i = next(k for k, p in enumerate(ps) if p.style.name.startswith("Heading 1") and name.lower() in p.text.lower())
    except StopIteration:
        return []
    out = []
    for p in ps[i + 1:]:
        if p.style.name.startswith("Heading 1"):
            break
        out.append(p)
    return out


def grade(run_dir, ev):
    res = []
    def A(text, ok, evidence):
        res.append({"text": text, "passed": bool(ok), "evidence": evidence})
    docs = [f for f in glob.glob(os.path.join(run_dir, "outputs", "**", "*.docx"), recursive=True) if "~$" not in f]
    A("Výstupem je soubor .docx", docs, ", ".join(os.path.basename(x) for x in docs) or "žádný docx")
    if not docs:
        return res
    f = max(docs, key=os.path.getmtime)
    d = Document(f)
    H = h1s(d)
    idx = []
    for s in POVINNE:
        m = [i for i, h in enumerate(H) if s.lower() in h.lower()]
        idx.append(m[0] if m else None)
    ok = None not in idx and idx == sorted(idx)
    A("Všech 7 povinných částí jako nadpisy 1. úrovně ve správném pořadí", ok, " | ".join(H)[:300])

    hdr = all("Mlna" in s.header._element.xml and "PAGE" in s.header._element.xml for s in d.sections)
    A("Záhlaví se jménem Matyáš Mlnařík a číslem stránky na všech stranách", hdr,
      f"{len(d.sections)} sekcí")

    body = d.element.body
    first_tbl_txt = d.tables[0].rows[0].cells[0].text if d.tables else ""
    has_desky = any("Vyšší odborná škola" in c.text for t in d.tables[:1] for r in t.rows[:2] for c in r.cells)
    A("Negeneruje desky (titulní list)", not has_desky, "nalezena tabulka desek" if has_desky else "bez desek")

    # grafy na samostatných stranách na šířku
    sekce, akt = [], []
    for el in body.iterchildren():
        akt.append(el)
        if el.tag == qn("w:p") and el.find(".//" + qn("w:sectPr")) is not None:
            sekce.append(akt); akt = []
    sekce.append(akt)
    graf_pages = 0
    for s, obs in zip(d.sections, sekce):
        if s.orientation == 1:
            n = sum(len(e.findall(".//" + qn("w:drawing"))) for e in obs)
            if n == 1:
                graf_pages += 1
    A(f"Alespoň {MIN_GRAFU.get(ev, 2)} grafů, každý sám na straně A4 na šířku", graf_pages >= MIN_GRAFU.get(ev, 2),
      f"{graf_pages} stran na šířku s jedním obrázkem")

    n_eq = len(body.findall(".//" + qn("m:oMath")))
    A("Rovnice jsou nativní rovnice Wordu (≥ 8)", n_eq >= 8, f"{n_eq} OMML rovnic")

    prist = [t for t in d.tables if "OZNA" in t.rows[0].cells[0].text.upper()]
    A("Tabulka přístrojů s 5 sloupci dle Zásad", prist and len(prist[0].columns) == 5,
      f"{len(prist[0].columns) if prist else 0} sloupců")

    # chybějící evidenční čísla nevymyšlená
    if prist:
        t = prist[0]
        rows = [[c.text.strip() for c in r.cells] for r in t.rows[1:]]
        hledej = {"stabilizatory-cv16": ["Uz1"], "usmernovac-cv14-chybejici-udaje": ["OSC", "A"],
                  "feromagnetika-cv11": []}.get(ev, [])
        bad = []
        for lab in hledej:
            r = next((r for r in rows if r[0].strip().upper() == lab.upper()), None)
            if r is None:
                bad.append(f"{lab}: řádek chybí"); continue
            ev_c = r[3] if len(r) > 3 else ""
            if re.search(r"\d{2,}", ev_c) and not re.search(r"doplnit|nemá|bez|\?|–|-$", ev_c, re.I):
                bad.append(f"{lab}: '{ev_c}'")
        if hledej:
            A("Chybějící evidenční čísla nejsou vymyšlená (označena k doplnění)", not bad, "; ".join(bad) or "ok")

    zav = section_texts(d, "Závěr")
    txt = " ".join(p.text for p in zav)
    lists = [p for p in zav if "List" in p.style.name or re.match(r"^\s*([-•–]|\d+[.)])\s", p.text)]
    zav_tbl = False
    A("Závěr jen v souvislých větách (bez odrážek)", zav and not lists, f"{len(lists)} odrážek")
    A("Závěr ≥ 120 slov a obsahuje konkrétní čísla", len(txt.split()) >= 120 and re.search(r"\d", txt),
      f"{len(txt.split())} slov")
    A("Závěr bez prázdných frází („dle očekávání“…)", not any(z in txt.lower() for z in ZAKAZ), "")

    tu = section_texts(d, "Teoretický úvod")
    tut = " ".join(p.text for p in tu)
    A("Teoretický úvod ≥ 500 slov", len(tut.split()) >= 500, f"{len(tut.split())} slov")
    A("Teoretický úvod uvádí literaturu/zdroje", re.search(r"literatur|zdroj", tut, re.I), "")

    vyp = section_texts(d, "Příklad výpočtu")
    bold_eq = 0
    for p in vyp:
        for om in p._p.findall(".//" + qn("m:oMath")):
            if om.find(".//" + qn("m:sty")) is not None and any(
                    s.get(qn("m:val")) in ("b", "bi") for s in om.iter(qn("m:sty"))):
                bold_eq += 1
    A("Výpočty mají tučně zvýrazněný výsledek", bold_eq >= 2, f"{bold_eq} tučných rovnic")

    alltext = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    import subprocess
    md = subprocess.run(["pandoc", f, "-t", "markdown"], capture_output=True, text=True).stdout
    if ev == "usmernovac-cv14-chybejici-udaje":
        A("Spočítán C0 = 3 mF (3000 µF)", re.search(r"3\s*mF|3000|3\s?000", md), "")
        A("Oscilogramy ponechány k doplnění (ne vymyšlené)", re.search(r"oscilogram", alltext, re.I) and re.search(r"DOPLNIT|doplní|dodá", alltext, re.I), "")
    if ev == "feromagnetika-cv11":
        A("Indukce počítána z U2 vztahem 4,44·f·N·S", re.search(r"4[,.]44", md), "")
        A("Uveden měrný ztrátový výkon / ztrátové číslo p1,0 nebo p1,5", re.search(r"p_?\{?1[,.][05]", md) or re.search(r"ztrátov\w+ čísl", alltext, re.I), "")
    if ev == "stabilizatory-cv16":
        A("Vypočten činitel stabilizace", re.search(r"činitel\w* stabilizace", alltext, re.I), "")
        A("Vypočten vnitřní odpor Ri obou stabilizátorů", len(re.findall(r"R_\{?i\}?\s*=", md)) >= 2, "")
    return res


if __name__ == "__main__":
    run_dir, ev = sys.argv[1], sys.argv[2]
    r = grade(run_dir, ev)
    passed = sum(x["passed"] for x in r)
    out = {"expectations": r, "summary": {"passed": passed, "failed": len(r) - passed, "total": len(r),
                                         "pass_rate": round(passed / len(r), 3) if r else 0}}
    json.dump(out, open(os.path.join(run_dir, "grading.json"), "w"), ensure_ascii=False, indent=1)
    for x in r:
        print("✔" if x["passed"] else "✘", x["text"], "–", x["evidence"])
